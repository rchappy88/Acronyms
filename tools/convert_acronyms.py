"""Convert source DOCX paragraphs into searchable JSON; retain source for audit.
Run: python tools/convert_acronyms.py assets/acronyms_mod.docx assets/acronyms.json
"""
import json,re,sys
from pathlib import Path
from docx import Document

def clean(s): return re.sub(r'\s+', ' ', s).strip()

def parse(parts):
    if len(parts)<2:return []
    out=[]; short=clean(parts[0]);
    for i, part in enumerate(parts[1:]):
        segment=clean(part)
        if i==len(parts)-2:
            if short and segment:out.append((short,segment))
            break
        # Mirrors original app's split heuristic; inspect flagged rows before release.
        match=re.search(r"\s+([A-Z0-9][A-Z0-9&/().*\-']{1,})$",segment)
        if match:
            meaning=clean(segment[:match.start()])
            if short and meaning:out.append((short,meaning))
            short=match.group(1)
    return out

def convert(src,dst):
    doc=Document(src); entries=[]; seen=set(); suspicious=[]
    for n,p in enumerate(doc.paragraphs,1):
        parts=p.text.split('\t'); rows=parse(parts)
        if len(parts)<2 or not rows:
            if clean(p.text): suspicious.append({'paragraph':n,'text':p.text[:220],'reason':'unparsed'})
            continue
        if len(parts)>2:suspicious.append({'paragraph':n,'text':p.text[:220],'reason':'multiple tabs / heuristic split'})
        for short,meaning in rows:
            key=(short.casefold(),meaning.casefold())
            if key in seen:continue
            seen.add(key)
            entries.append({'shortName':short,'meaning':meaning,'category':'MOD directory','description':''})
    Path(dst).parent.mkdir(parents=True,exist_ok=True)
    Path(dst).write_text(json.dumps(entries,ensure_ascii=False,separators=(',',':')),encoding='utf-8')
    report={'source_paragraphs':len(doc.paragraphs),'unique_entries':len(entries),'review_rows':len(suspicious),'review_samples':suspicious[:150]}
    Path(dst).with_suffix('.audit.json').write_text(json.dumps(report,indent=2,ensure_ascii=False),encoding='utf-8')
    print(f"Exported {len(entries)} entries; {len(suspicious)} paragraphs flagged for review")
if __name__=='__main__':convert(sys.argv[1],sys.argv[2])
